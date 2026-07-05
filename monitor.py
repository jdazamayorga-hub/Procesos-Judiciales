# monitor.py — Monitor diario de procesos judiciales (Rama Judicial CO)
import csv, json, time, os, sys, unicodedata, smtplib, datetime
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import urllib.request, urllib.error, ssl
import config

BASE = "https://consultaprocesos.ramajudicial.gov.co:448/api/v2"
CARPETA = os.path.dirname(os.path.abspath(__file__))
F_EXP = os.path.join(CARPETA, "expedientes.csv")
F_EST = os.path.join(CARPETA, "estado.json")
F_HTML = os.path.join(CARPETA, "reporte.html")
CTX = ssl.create_default_context()

def sin_tildes(t):
    t = unicodedata.normalize("NFD", t)
    return "".join(c for c in t if unicodedata.category(c) != "Mn").lower()

def pedir(url):
    for i in range(config.REINTENTOS + 1):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=config.TIMEOUT, context=CTX) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception as e:
            if i == config.REINTENTOS:
                raise
            time.sleep(config.PAUSA_ENTRE_CONSULTAS)
    return None

def id_proceso(radicado):
    url = f"{BASE}/Procesos/Consulta/NumeroRadicacion?numero={radicado}&SoloActivos=false&pagina=1"
    d = pedir(url)
    procesos = d.get("procesos") or []
    if not procesos:
        return None
    return procesos[0].get("idProceso")

def ultima_actuacion(idp):
    url = f"{BASE}/Proceso/Actuaciones/{idp}"
    d = pedir(url)
    acts = d.get("actuaciones") or []
    if not acts:
        return None
    a = acts[0]  # la API devuelve la más reciente primero
    return {
        "fecha": a.get("fechaActuacion", "") or a.get("fechaRegistro", ""),
        "actuacion": a.get("actuacion", ""),
        "anotacion": a.get("anotacion", "") or "",
    }

def criticidad(texto):
    t = sin_tildes(texto)
    if any(sin_tildes(k) in t for k in config.CRITICIDAD_ALTA):
        return "ALTA"
    if any(sin_tildes(k) in t for k in config.CRITICIDAD_MEDIA):
        return "MEDIA"
    return "BAJA"

def cargar_estado():
    if os.path.exists(F_EST):
        try:
            with open(F_EST, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def guardar_estado(e):
    with open(F_EST, "w", encoding="utf-8") as f:
        json.dump(e, f, ensure_ascii=False, indent=2)

def leer_expedientes():
    filas = []
    with open(F_EXP, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            rad = (row.get("radicado") or "").strip()
            if rad:
                filas.append({"radicado": rad, "alias": (row.get("alias") or "").strip()})
    return filas

def avisar_error(errores):
    if not config.CORREO_ACTIVO or not errores:
        return
    cuerpo = "El monitor judicial encontró problemas hoy:\n\n" + "\n".join(errores)
    msg = MIMEMultipart()
    msg["From"] = config.CORREO_REMITENTE
    msg["To"] = config.CORREO_DESTINO
    msg["Subject"] = f"[Monitor Judicial] Errores {datetime.date.today()}"
    msg.attach(MIMEText(cuerpo, "plain"))
    try:
        s = smtplib.SMTP(config.SMTP_SERVIDOR, config.SMTP_PUERTO)
        s.starttls()
        s.login(config.CORREO_REMITENTE, config.CORREO_CLAVE_APP)
        s.send_message(msg)
        s.quit()
    except Exception as e:
        print("No se pudo enviar correo:", e)

def color(c):
    return {"ALTA": "#d32f2f", "MEDIA": "#ef6c00", "BAJA": "#616161"}[c]

def generar_html(resultados, errores):
    hoy = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    con_nov = [r for r in resultados if r["novedad"]]
    filas = ""
    orden = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
    for r in sorted(resultados, key=lambda x: (not x["novedad"], orden.get(x["crit"], 3))):
        nov = "🔴 SÍ" if r["novedad"] else "—"
        peso = "font-weight:700;" if r["novedad"] else ""
        link = f'https://consultaprocesos.ramajudicial.gov.co/Procesos/NumeroRadicacion'
        filas += f"""<tr style="{peso}">
<td>{r['alias']}</td>
<td style="font-family:monospace;font-size:12px">{r['radicado']}</td>
<td style="text-align:center">{nov}</td>
<td><span style="background:{color(r['crit'])};color:#fff;padding:2px 8px;border-radius:4px;font-size:12px">{r['crit']}</span></td>
<td>{r['fecha']}</td>
<td>{r['actuacion']}<br><span style="color:#555;font-size:12px">{r['anotacion'][:200]}</span></td>
</tr>"""
    err_html = ""
    if errores:
        err_html = '<div style="background:#fff3e0;border-left:4px solid #ef6c00;padding:12px;margin:16px 0"><b>Problemas en la consulta:</b><br>' + "<br>".join(errores) + "</div>"
    html = f"""<!DOCTYPE html><html lang="es"><head><meta charset="utf-8">
<title>Monitor Judicial — {hoy}</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#212121}}
h1{{font-size:20px}} .sub{{color:#666;font-size:13px;margin-bottom:16px}}
table{{border-collapse:collapse;width:100%;font-size:14px}}
th{{background:#263238;color:#fff;padding:8px;text-align:left}}
td{{padding:8px;border-bottom:1px solid #e0e0e0;vertical-align:top}}
tr:hover{{background:#f5f5f5}}</style></head><body>
<h1>Monitor de procesos judiciales</h1>
<div class="sub">Generado: {hoy} · Expedientes: {len(resultados)} · Con novedad hoy: {len(con_nov)}</div>
{err_html}
<table><thead><tr><th>Expediente</th><th>Radicado</th><th>Novedad</th><th>Criticidad</th><th>Fecha última actuación</th><th>Detalle</th></tr></thead>
<tbody>{filas}</tbody></table>
<p class="sub">Verifique siempre la actuación en el portal oficial antes de decidir. Este reporte es una ayuda, no reemplaza la consulta directa.</p>
</body></html>"""
    with open(F_HTML, "w", encoding="utf-8") as f:
        f.write(html)

def generar_excel(con_nov):
    if not con_nov:
        return None
    try:
        import openpyxl
    except ImportError:
        os.system(f'"{sys.executable}" -m pip install openpyxl')
        import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Novedades"
    cab = ["Expediente", "Radicado", "Criticidad", "Fecha", "Actuación", "Anotación"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="263238")
    for r in con_nov:
        ws.append([r["alias"], r["radicado"], r["crit"], r["fecha"], r["actuacion"], r["anotacion"]])
    for r in ws.iter_rows(min_row=2):
        if r[2].value == "ALTA":
            r[2].fill = PatternFill("solid", fgColor="FFCDD2")
    anchos = [28, 26, 12, 20, 30, 60]
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    nombre = os.path.join(CARPETA, f"novedades_{datetime.date.today()}.xlsx")
    wb.save(nombre)
    return nombre

def main():
    exps = leer_expedientes()
    estado = cargar_estado()
    resultados, errores = [], []
    for e in exps:
        rad, alias = e["radicado"], e["alias"] or e["radicado"]
        try:
            idp = id_proceso(rad)
            if not idp:
                errores.append(f"{alias} ({rad}): sin resultados o radicado no encontrado")
                resultados.append({"alias": alias, "radicado": rad, "novedad": False,
                                   "crit": "BAJA", "fecha": "?", "actuacion": "NO ENCONTRADO", "anotacion": ""})
                continue
            act = ultima_actuacion(idp)
            time.sleep(config.PAUSA_ENTRE_CONSULTAS)
            if not act:
                errores.append(f"{alias} ({rad}): proceso sin actuaciones registradas")
                continue
            firma = f"{act['fecha']}|{act['actuacion']}"
            prev = estado.get(rad, {}).get("firma")
            novedad = (prev is not None and prev != firma)
            texto = f"{act['actuacion']} {act['anotacion']}"
            resultados.append({"alias": alias, "radicado": rad, "novedad": novedad,
                               "crit": criticidad(texto), "fecha": act["fecha"],
                               "actuacion": act["actuacion"], "anotacion": act["anotacion"]})
            estado[rad] = {"firma": firma, "revisado": str(datetime.datetime.now())}
        except Exception as ex:
            errores.append(f"{alias} ({rad}): error de consulta — {type(ex).__name__}: {ex}")
        time.sleep(config.PAUSA_ENTRE_CONSULTAS)

    guardar_estado(estado)
    generar_html(resultados, errores)
    con_nov = [r for r in resultados if r["novedad"]]
    xls = generar_excel(con_nov)
    avisar_error(errores)

    print(f"Listo. {len(resultados)} expedientes. {len(con_nov)} con novedad.")
    print(f"Reporte: {F_HTML}")
    if xls: print(f"Excel: {xls}")
    if errores: print(f"Errores: {len(errores)}")
    try:
        os.startfile(F_HTML)   # abre el reporte solo (Windows)
    except Exception:
        pass

if __name__ == "__main__":
    main()
